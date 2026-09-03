import logging
import csv
import os
from dataclasses import dataclass
from typing import List, Optional, Dict, Any

logger = logging.getLogger(__name__)

@dataclass
class ParsedSection:
    content: str
    heading_path: Optional[str]
    page_number: Optional[int]
    content_type: str  # 'text' | 'table' | 'list'
    metadata: Dict[str, Any]

def table_to_markdown(table: List[List[Optional[str]]]) -> str:
    """
    Converts a 2D list (table) into a Markdown formatted table string.
    Handles None cells, empty rows, and malformed tables gracefully.
    """
    if not table:
        return ""
    
    cleaned_table = []
    max_cols = 0
    for row in table:
        if not row:
            continue
        if not isinstance(row, list):
            row = list(row)
        cleaned_row = [str(cell).strip().replace('\n', ' ') if cell is not None else "" for cell in row]
        if any(cleaned_row):
            cleaned_table.append(cleaned_row)
            max_cols = max(max_cols, len(cleaned_row))

    if not cleaned_table:
        return ""

    for row in cleaned_table:
        while len(row) < max_cols:
            row.append("")

    md_lines = []
    headers = cleaned_table[0]
    md_lines.append("| " + " | ".join(headers) + " |")
    md_lines.append("|" + "|".join(["---"] * max_cols) + "|")
    
    for row in cleaned_table[1:]:
        md_lines.append("| " + " | ".join(row) + " |")
        
    return "\n".join(md_lines)

import re

# Pattern-based heuristic for detecting section headings in generated PDFs.
# Matches lines like "4. Data Subject Rights", "20. Contact Information", "4.1 Scope".
# NOTE: This is a format-specific heuristic, not general PDF structure
# parsing. It works for the current corpus (numbered section headings)
# but will NOT detect arbitrary headings in PDFs that use different formatting
# (e.g. bold text without numbering, custom fonts, etc.).
_PDF_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2})*)\.?\s+([A-Z][A-Za-z0-9\s,\-\(\)\/&]+)$")


def parse_pdf(file_path: str) -> List[ParsedSection]:
    """Parse PDF file into text and table sections with pattern-based heading detection."""
    sections = []
    try:
        import pdfplumber
    except ImportError:
        logger.error("pdfplumber not installed. Cannot parse PDF.")
        return sections

    try:
        current_heading: Optional[str] = None

        with pdfplumber.open(file_path) as pdf:
            for page_idx, page in enumerate(pdf.pages):
                page_num = page_idx + 1

                text = page.extract_text()
                if text:
                    lines = text.splitlines()
                    current_section_lines = []

                    for line in lines:
                        trimmed = line.strip()
                        if not trimmed:
                            continue

                        # Check if line matches a numbered section heading
                        if _PDF_HEADING_RE.match(trimmed) and len(trimmed) < 100:
                            # Flush accumulated text for previous heading if any
                            if current_section_lines:
                                sections.append(
                                    ParsedSection(
                                        content="\n".join(current_section_lines),
                                        heading_path=current_heading,
                                        page_number=page_num,
                                        content_type="text",
                                        metadata={},
                                    )
                                )
                                current_section_lines = []
                            current_heading = trimmed

                        current_section_lines.append(line)

                    # Flush remaining text on this page
                    if current_section_lines:
                        sections.append(
                            ParsedSection(
                                content="\n".join(current_section_lines),
                                heading_path=current_heading,
                                page_number=page_num,
                                content_type="text",
                                metadata={},
                            )
                        )

                tables = page.extract_tables()
                for table_idx, table in enumerate(tables):
                    md_table = table_to_markdown(table)
                    if md_table:
                        sections.append(
                            ParsedSection(
                                content=md_table,
                                heading_path=current_heading,
                                page_number=page_num,
                                content_type="table",
                                metadata={"table_index": table_idx},
                            )
                        )

            if not sections:
                logger.warning(
                    f"No text or tables found in PDF {file_path}. Might be a scanned image."
                )
    except Exception as e:
        logger.error(f"Error parsing PDF {file_path}: {e}")

    return sections

def parse_docx(file_path: str) -> List[ParsedSection]:
    """Parse DOCX file into text (with headings) and table sections."""
    sections = []
    try:
        from docx import Document
    except ImportError:
        logger.error("python-docx not installed. Cannot parse DOCX.")
        return sections

    try:
        doc = Document(file_path)
        current_headings = {}

        def get_heading_path() -> Optional[str]:
            if not current_headings:
                return None
            levels = sorted(current_headings.keys())
            return " > ".join(current_headings[lvl] for lvl in levels)

        # Real Word headers/footers live in a separate part of the document
        # structure — doc.paragraphs (used below) does NOT include them at
        # all, regardless of how much of the body you scan. Extract them
        # explicitly and place header content first, footer content last,
        # so downstream first-N/last-N logic naturally picks them up.
        header_sections: List[ParsedSection] = []
        footer_sections: List[ParsedSection] = []
        _seen_headers: set[str] = set()
        _seen_footers: set[str] = set()
        for sec in doc.sections:
            for hp in sec.header.paragraphs:
                text = hp.text.strip()
                if text and text not in _seen_headers:
                    _seen_headers.add(text)
                    header_sections.append(ParsedSection(
                        content=text, heading_path="[Header]",
                        page_number=None, content_type='text', metadata={'source': 'header'}
                    ))
            for fp in sec.footer.paragraphs:
                text = fp.text.strip()
                if text and text not in _seen_footers:
                    _seen_footers.add(text)
                    footer_sections.append(ParsedSection(
                        content=text, heading_path="[Footer]",
                        page_number=None, content_type='text', metadata={'source': 'footer'}
                    ))

        body_sections: List[ParsedSection] = []

        for paragraph in doc.paragraphs:
            style_name = paragraph.style.name if paragraph.style else ""
            text = paragraph.text.strip()
            
            if not text:
                continue
                
            if style_name.startswith('Heading'):
                try:
                    level = int(style_name.split(' ')[-1])
                    keys_to_remove = [k for k in current_headings.keys() if k >= level]
                    for k in keys_to_remove:
                        del current_headings[k]
                    current_headings[level] = text
                except ValueError:
                    pass
            
            body_sections.append(ParsedSection(
                content=text,
                heading_path=get_heading_path(),
                page_number=None,
                content_type='text',
                metadata={'style': style_name}
            ))
            
        for table_idx, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                table_data.append([cell.text for cell in row.cells])
            
            md_table = table_to_markdown(table_data)
            if md_table:
                body_sections.append(ParsedSection(
                    content=md_table,
                    heading_path=get_heading_path(),
                    page_number=None,
                    content_type='table',
                    metadata={'table_index': table_idx}
                ))

        # Header first, footer last — this ordering is what lets main.py's
        # first-N/last-N aggregation naturally capture header/footer
        # content without any special-casing on the caller's side.
        sections = header_sections + body_sections + footer_sections

    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path}: {e}")
        
    return sections

def parse_xlsx(file_path: str) -> List[ParsedSection]:
    """Parse XLSX file sheets into table sections."""
    sections = []
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl not installed. Cannot parse XLSX.")
        return sections

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            table_data = []
            for row in sheet.iter_rows(values_only=True):
                table_data.append(list(row))
            
            md_table = table_to_markdown(table_data)
            if md_table:
                sections.append(ParsedSection(
                    content=md_table,
                    heading_path=sheet_name,
                    page_number=None,
                    content_type='table',
                    metadata={'sheet_name': sheet_name}
                ))
    except Exception as e:
        logger.error(f"Error parsing XLSX {file_path}: {e}")
        
    return sections

def parse_csv(file_path: str) -> List[ParsedSection]:
    """Parse CSV file into a table section."""
    sections = []
    try:
        with open(file_path, 'r', encoding='utf-8', newline='') as f:
            reader = csv.reader(f)
            table_data = list(reader)
            
            md_table = table_to_markdown(table_data)
            if md_table:
                sections.append(ParsedSection(
                    content=md_table,
                    heading_path=None,
                    page_number=None,
                    content_type='table',
                    metadata={}
                ))
    except Exception as e:
        logger.error(f"Error parsing CSV {file_path}: {e}")
        
    return sections

def parse_txt(file_path: str) -> List[ParsedSection]:
    """Parse plain text or markdown file into a text section."""
    sections = []
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
            if content.strip():
                sections.append(ParsedSection(
                    content=content,
                    heading_path=None,
                    page_number=None,
                    content_type='text',
                    metadata={}
                ))
    except Exception as e:
        logger.error(f"Error parsing TXT {file_path}: {e}")
        
    return sections

def parse_file(file_path: str, filename: str) -> List[ParsedSection]:
    """
    Main entry point for parsing files into sections.
    Dispatches to the correct parser based on file extension.
    """
    ext = os.path.splitext(filename.lower())[1]
    
    if ext == '.pdf':
        return parse_pdf(file_path)
    elif ext == '.docx':
        return parse_docx(file_path)
    elif ext == '.xlsx':
        return parse_xlsx(file_path)
    elif ext == '.csv':
        return parse_csv(file_path)
    elif ext in ['.txt', '.md']:
        return parse_txt(file_path)
    else:
        logger.warning(f"Unsupported file extension: {ext} for {filename}")
        return []
